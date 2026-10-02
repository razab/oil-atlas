"""Strict, offline extraction of the pinned annual source boundaries.
Numbers retain original units; no density-based mass/volume conversion is invented.
"""
import hashlib, json, math, re
from html import unescape
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
RAW=ROOT/'data/raw'
SUBTOTALS={'Europe','North America','S. & Cent. America','Middle East & Africa','CIS','Middle East','Africa','Asia Pacific'}
FR_NAMES={'Algérie':'DZA','Libye':'LBY','Nigéria':'NGA','dont États-Unis':'USA','Kazakhstan':'KAZ','Russie':'RUS','dont Norvège':'NOR','dont Arabie saoudite':'SAU'}

def clean(s): return ' '.join(unescape(re.sub('<[^>]+>',' ',s)).split())
def numeric(v):
    if v is None or str(v).strip() in {'','--','NA','N/A','(s)','W','-'}:return None
    if isinstance(v,(float,int)):
        if math.isfinite(v) and v>=0:return float(v)
        raise ValueError(f'Invalid source number: {v}')
    s=str(v).strip().replace(',','')
    if re.fullmatch(r'(\d+(\.\d+)?|\.\d+)',s):return float(s)
    raise ValueError(f'Unknown source flag: {v!r}')
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def aliases():
    countries=json.loads((RAW/'countries.json').read_text())
    result={name:c['cca3'] for c in countries for name in [c['name']['common'],c['name']['official']]+c.get('altSpellings',[])}
    result.update({'US':'USA','U.S.':'USA','United States':'USA','Russian Federation':'RUS','Russia':'RUS','IR Iran':'IRN','Iran':'IRN','Türkiye':'TUR','Turkiye':'TUR','Turkey':'TUR','South Korea':'KOR','Korea, South':'KOR','Trinidad & Tobago':'TTO','Taiwan':'TWN','Congo':'COG','Republic of Congo':'COG','Congo (Brazzaville)':'COG','Congo (Kinshasa)':'COD','Bahama Islands':'BHS','Ivory Coast':'CIV','Virgin Islands (U.S.)':'VIR','Virgin Islands (British)':'VGB','St. Lucia':'LCA','St. Vincent and the Grenadines':'VCT','Georgia, Republic of':'GEO','Gibralter':'GIB','Djbouti':'DJI','Macau S.A.R.':'MAC','Micronesia, Federated States of':'FSM','Serbia (Excludes Kosovo)':'SRB','China Hong Kong SAR':'HKG','Czech Republic':'CZE','UAE':'ARE'})
    return result

def france_rows(path):
    h=path.read_text();table=re.search(r'<table id="produit-tableau-figure1".*?</table>',h,re.S)
    if not table or 'en millions de tonnes' not in table[0]:raise ValueError('France: unit/table boundary changed')
    headers=[clean(s) for s in re.findall(r'<th[^>]*scope="col"[^>]*>(.*?)</th>',table[0],re.S)]
    if headers.count('2025 (p)')!=1:raise ValueError('France: annual mass column is ambiguous')
    col=headers.index('2025 (p)')-1;rows=[];countries={};total=None
    for i,r in enumerate(re.findall(r'<tr.*?</tr>',table[0],re.S),1):
        label=re.search(r'<th[^>]*scope="row"[^>]*>(.*?)</th>',r,re.S)
        if not label:continue
        label=clean(label[1]);cells=re.findall(r'<td[^>]*>(.*?)</td>',r,re.S)
        if len(cells)!=len(headers)-1:
            if label=='dont :':continue
            raise ValueError(f'France: unexpected cells in {label}')
        v=numeric(clean(cells[col]).replace(',','.'))
        rows.append(dict(label=label,value=v,locator=f'produit-tableau-figure1/row={i}/column={headers[col+1]}'))
        if label in FR_NAMES:countries[FR_NAMES[label]]=v
        if label=='Importations totales':total=v
    if set(countries)!=set(FR_NAMES.values()) or total is None:raise ValueError('France: missing supplier or total')
    return dict(countries=countries,total=total,rows=rows)

def eia_rows(path):
    h=path.read_text();years=re.findall(r'<th class="Series5">(\d{4})</th>',h)
    if not years or years[-1]!='2025' or years.count('2025')!=1:raise ValueError(f'EIA: wrong annual column {path.name}')
    if 'Thousand Barrels' not in clean(h):raise ValueError(f'EIA: wrong unit {path.name}')
    result=[]
    for i,r in enumerate(re.split(r'<tr class="DataRow">',h)[1:],1):
        n=re.search(r'<td class="DataStub1">(.*?)</td>',r,re.S)
        cells=re.findall(r'<td[^>]*class="(?:DataB?|Current2)"[^>]*>(.*?)</td>',r,re.S)
        if not n:continue
        # Every row has one value per exact year header; current cell is the last annual column.
        if len(cells)!=len(years):raise ValueError(f'EIA: unexpected year cells {path.name}, row {i}: {len(cells)}')
        raw=clean(cells[years.index('2025')]);v=numeric(raw)
        result.append(dict(label=clean(n[1]),value=v,raw=raw,locator=f'DataRow={i}/year=2025'))
    if not result:raise ValueError('EIA: no annual rows')
    return result

def extract(scope):
    import openpyxl
    registry=json.loads((ROOT/'scripts/source-lock.json').read_text());source_files={}
    def use(file):
        actual=sha(RAW/file)
        if registry['files'].get(file)!=actual:raise ValueError(f'Unreviewed source edition: {file}')
        source_files[file]=dict(sha256=actual,**registry['metadata'].get(file,{}))
    a=aliases();facts=[];checks=[];excluded=[]
    def fact(source,locator,value,unit,year,**kw):
        f=dict(id=source+'#'+locator,source=source,locator=locator,value=value,unit=unit,year=year,**kw);facts.append(f);return f
    def check(label,actual,expected,tolerance,unit):
        checks.append(dict(label=label,actual=actual,expected=expected,tolerance=tolerance,unit=unit))
        if not math.isclose(actual,expected,abs_tol=tolerance,rel_tol=0):raise ValueError(f'{label}: {actual} != {expected} (±{tolerance})')
    use('countries.json');use('EI-2026.xlsx')
    w=openpyxl.load_workbook(RAW/'EI-2026.xlsx',read_only=True,data_only=True)
    for sheet,measure,product in [('Oil Production - barrels','production','oil_liquids'),('Oil Consumption - barrels','consumption','oil_liquids'),('Gas Production - Bcm','production','natural_gas'),('Gas Consumption - Bcm','consumption','natural_gas')]:
        rows=list(w[sheet].values);col=list(rows[2]).index(2025)
        if rows[2][col-1]!=2024:raise ValueError('Not the annual year column')
        unit='kbpd' if product=='oil_liquids' else 'bcm'
        if rows[2][0]!=('Thousand barrels daily' if unit=='kbpd' else 'Billion cubic metres'):raise ValueError('Wrong EI unit')
        if scope=='full':
            regional=[r for r in rows[4:] if str(r[0] or '').strip() in {'Total North America','Total S. & Cent. America','Total Europe','Total CIS','Total Middle East','Total Africa','Total Asia Pacific'}]
            world=next(r[col] for r in rows if r[0]=='Total World')
            if len(regional)!=7:raise ValueError('EI annual regional boundary changed')
            check(sheet+' / seven regions vs world',sum(r[col] for r in regional),world,1e-7,unit)
            # Each regional block includes named countries and residual groups exactly once.
            block=[]
            for r in rows[4:]:
                name=str(r[0] or '').strip()
                if name=='Total World':break
                if not name:continue
                if name.startswith('Total '):
                    check(sheet+' / '+name,sum(numeric(x[col]) or 0 for x in block),r[col],1e-7,unit);block=[]
                else:block.append(r)
        basis=('crude, condensates, oil sands and NGL; excludes biofuels, synthetic fuels and refinery gain' if measure=='production' else 'inland demand, aviation/marine bunkers, refinery fuel/loss; excludes biofuels; includes coal/gas derivatives') if product=='oil_liquids' else 'annual natural gas; gas-equivalent volume standardized at GCV 40 MJ/m3'
        for i,row in enumerate(rows[4:],5):
            name=str(row[0] or '').strip();code=a.get(name)
            if not name:continue
            if scope=='canary' and code not in {'USA','FRA','RUS','LUX','JPN','CHN','SAU','CAN','MEX'} and name!='Total World':continue
            if code or name=='Total World':
                fact('EI-2026.xlsx',f'{sheet}!{openpyxl.utils.get_column_letter(col+1)}{i}',numeric(row[col]),unit,2025,country=code or 'WORLD',product=product,measure=measure,basis=basis)
            else:excluded.append(dict(source='EI-2026.xlsx',table=sheet,label=name,reason='Aggregate or footnote, not a country'))
    for sheet,product,header in [('Gas trade 2025 - LNG','lng',3),('Gas trade 2025 - pipeline','pipeline',2)]:
        rows=list(w[sheet].values);labels=[str(v or '').strip() for v in rows[header]]
        if '2025' not in rows[0][0] or rows[header-1][0]!='Billion cubic metres':raise ValueError('EI gas boundary changed')
        if product=='lng':labels=[' '.join([str(rows[2][i] or '').strip(),str(rows[3][i] or '').strip()]).strip() for i in range(len(labels))]
        totalcol=labels.index('Total imports');end=next(i for i,r in enumerate(rows) if r[0]=='Total exports')
        leaves=[(i,r) for i,r in enumerate(rows[header+1:end],header+1) if str(r[0]).strip() not in SUBTOTALS]
        code=lambda n:a.get(n) or 'region:'+n
        for i,row in leaves:
            dest=code(str(row[0]).strip());vs=[numeric(v) for v in row[1:totalcol]]
            if any(v is None for v in vs):raise ValueError('Missing EI gas cell requires explicit treatment')
            if scope=='full' or dest in {'USA','FRA','RUS'}:check(sheet+' / imports '+str(row[0]),sum(vs),row[totalcol],1e-8,'bcm')
            for j,v in enumerate(vs,1):
                origin=code(labels[j])
                if scope=='canary' and not ({origin,dest}&{'USA','FRA','RUS'}):continue
                fact('EI-2026.xlsx',f'{sheet}!{openpyxl.utils.get_column_letter(j+1)}{i+1}',v,'bcm',2025,origin=origin,destination=dest,product=product,measure='trade',reporter=None,basis='bilateral matrix; gas-equivalent volume; GCV 40 MJ/m3; re-exports included' if product=='lng' else 'bilateral matrix; intra-region trade excluded; GCV 40 MJ/m3')
        for j,n in enumerate(labels[1:totalcol],1):
            if scope=='full' or code(n) in {'RUS','USA'}:check(sheet+' / exports '+n,sum(r[j] for _,r in leaves),rows[end][j],1e-8,'bcm')
        if scope=='full':check(sheet+' / world',sum(r[totalcol] for _,r in leaves),rows[end][totalcol],1e-8,'bcm')
    for product in ['crude','gasoline','diesel']:
        for direction in ['imports','exports']:
            fn=f'us-{product}-{direction}-2025.html';use(fn);rows=eia_rows(RAW/fn)
            total=next(r for r in rows if r['label'] in {'All Countries','Total All Countries'})
            leaves=[r for r in rows if r['label'] not in {'All Countries','Total All Countries','Persian Gulf','OPEC*','Non OPEC*'}]
            numericrows=[r for r in leaves if r['value'] is not None]
            check(fn+' / countries vs total',sum(r['value'] for r in numericrows),total['value'],.5*(len(numericrows)+1),'thousand_barrels/year')
            fact(fn,total['locator'],total['value'],'thousand_barrels/year',2025,country='USA',product=product,measure='total_'+direction,basis='EIA published national trade total')
            for r in leaves:
                partner=a.get(r['label'])
                if not partner:
                    if r['value'] and r['value']>0:raise ValueError(f'Unmapped positive EIA country: {r["label"]}')
                    excluded.append(dict(source=fn,label=r['label'],value=r['value'],reason='Historical or unspecified category; no positive 2025 observation'));continue
                origin,dest=(partner,'USA') if direction=='imports' else ('USA',partner)
                fact(fn,r['locator'],r['value'],'thousand_barrels/year',2025,origin=origin,destination=dest,product=product,measure='trade',reporter='USA',rawText=r['raw'],basis='EIA customs import origin' if direction=='imports' else 'EIA export destination')
    fn='france-crude-imports-2025.html';use(fn);fr=france_rows(RAW/fn)
    regions=[r for r in fr['rows'] if r['label'] in {'Afrique','Amérique du Nord','URSS/ex-URSS','Mer du Nord','Moyen-Orient','Autres'}]
    check('France / regions vs total',sum(r['value'] for r in regions),fr['total'],.05*(len(regions)+1),'Mt')
    for r in fr['rows']:
        country=FR_NAMES.get(r['label'])
        if country:fact(fn,r['locator'],r['value'],'Mt',2025,origin=country,destination='FRA',product='crude',measure='trade',reporter='FRA',basis='country of extraction; condensates and other refinery feedstocks included; provisional')
        elif r['label']=='Importations totales':fact(fn,r['locator'],r['value'],'Mt',2025,country='FRA',product='crude',measure='total_imports',basis='country of extraction; provisional')
        else:excluded.append(dict(source=fn,label=r['label'],value=r['value'],reason='Region or overlapping subtotal; not an additional country flow'))
    # State totals exclude nested Alaska subareas and PADD aggregates, include federal offshore.
    fn='us-state-crude-2025.html';use(fn);states={g['properties']['name'] for g in json.loads((ROOT/'data/us-states.json').read_text())['objects']['states']['geometries']}
    rows=eia_rows(RAW/fn);leaves=[r for r in rows if r['label'] in states or r['label'].startswith('Federal Offshore')]
    check('US / states + federal offshore',sum(r['value'] or 0 for r in leaves),next(r['value'] for r in rows if r['label']=='U.S.'),.5*(len(leaves)+1),'thousand_barrels/year')
    for r in leaves:fact(fn,r['locator'],r['value'],'thousand_barrels/year',2025,country='USA',subdivision=r['label'],product='crude',measure='production',basis='crude including lease condensate')
    fn='us-fuel-supply-2025.html';use(fn);h=(RAW/fn).read_text()
    if '2025 (Current)' not in h or 'Thousand Barrels' not in clean(h):raise ValueError('US supply year/unit changed')
    for key,product,measure in [('MCRFPUS1','crude','production'),('MCRRIUS1','crude','refinery_input'),('MGFRPUS1','gasoline','production'),('MGFUPUS1','gasoline','product_supplied'),('MDIRPUS1','diesel','production'),('MDIUPUS1','diesel','product_supplied')]:
        values=re.findall(r'&s='+re.escape(key)+r'&f=A[^>]*>([\d,]+)</a>',h)
        if len(values)!=1:raise ValueError('US supply duplicate or missing series')
        fact(fn,f'series={key}/year=2025',numeric(values[0]),'thousand_barrels/year',2025,country='USA',product=product,measure=measure,basis='EIA annual supply; refinery input is not final demand; distillates include heating oil')
    fn='us-basins-2025.html';use(fn);h=(RAW/fn).read_text()
    paragraphs=[clean(x) for x in re.findall(r'<p>(.*?)</p>',h,re.S)]
    patterns=[('permian',r'In 2025, the Permian region .*?to ([\d.]+) million b/d'),('eagle-ford',r'Eagle Ford production .*?to ([\d.]+) million b/d in 2025'),('bakken',r'Bakken production fell .*?to ([\d.]+) million b/d'),('gulf',r'Crude oil production in GOA .*?in 2025, averaging ([\d.]+) million b/d')]
    for site,pattern in patterns:
        hits=[(i,m) for i,p in enumerate(paragraphs,1) if (m:=re.search(pattern,p))]
        if len(hits)!=1:raise ValueError('Basin source boundary changed: '+site)
        i,m=hits[0];fact(fn,f'article-paragraph={i}/basin={site}',numeric(m[1]),'million_bbl/day',2025,country='USA',subdivision=site,product='crude',measure='production',basis='EIA STEO March2026; rounded regional estimate; overlaps state totals, do not add together')
    fn='opec-reserves-table-3.1-2024.txt';use(fn);h=(RAW/fn).read_text()
    if not all(s in h for s in ['Table 3.1','2024','(mb)','Data excludes oil sands']):raise ValueError('OPEC reserves boundary changed')
    for i,line in enumerate(h.splitlines(),1):
        m=re.fullmatch(r'(.+?)\s{2,}([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+[–\-\d,]+\s*',line)
        if not m:continue
        name=m[1].strip().removesuffix('1');code=a.get(name)
        if name=='Total world':code='WORLD'
        if not code:continue
        if scope=='canary' and code not in {'USA','CAN','RUS','JPN'}:continue
        fact(fn,f'table=3.1/line={i}/year=2024',numeric(m[6]),'million_barrels',2024,country=code,product='crude',measure='reserves',basis='proven crude reserves; year end; Canada excludes oil sands')
    if scope=='full':
        fn='eia-crude-ngl-2025.json';use(fn);response=json.loads((RAW/fn).read_text())['response']
        if len(response['data'])!=int(response['total']):raise ValueError('EIA response truncated')
        for i,r in enumerate(response['data']):
            if r['period']!='2025' or r['unit']!='TBPD' or r['activityId']!='1' or r['productId'] not in {'57','58'}:raise ValueError('EIA API boundary changed')
            code=r['countryRegionId']
            if r['countryRegionTypeId']!='c' or code not in a.values():
                excluded.append(dict(source=fn,label=code,reason='Aggregate or historical entity'));continue
            fact(fn,f'response.data[{i}]',numeric(r['value']),'kbpd',2025,country=code,product='crude' if r['productId']=='57' else 'ngpl',measure='production',rawText=str(r['value']),basis='EIA international annual crude/NGPL, separate products')
        use('EI-2025-trade.xlsx');old=openpyxl.load_workbook(RAW/'EI-2025-trade.xlsx',read_only=True,data_only=True);sheet='Oil - Inter-area movements';rows=list(old[sheet].values)
        if '2024' not in rows[0][0] or rows[2][0]!='Crude (million tonnes)':raise ValueError('EI oil year/unit changed')
        labels=rows[2];totalcol=labels.index('Total');end=next(i for i,r in enumerate(rows) if r[0]=='Total imports');leaves=list(enumerate(rows[4:end],4))
        code=lambda n:a.get(n) or 'region:'+n
        for i,r in leaves:
            vs=[numeric(v) for v in r[1:totalcol]]
            check('EI oil2024 / exports '+r[0],sum(v or 0 for v in vs),r[totalcol],1e-7,'Mt')
            for j,v in enumerate(vs,1):fact('EI-2025-trade.xlsx',f'{sheet}!{openpyxl.utils.get_column_letter(j+1)}{i+1}',v,'Mt',2024,origin=code(r[0]),destination=code(labels[j]),product='crude',measure='trade',reporter=None,basis='inter-area crude matrix; regional groups retained')
        for j,n in enumerate(labels[1:totalcol],1):check('EI oil2024 / imports '+n,sum(r[j] for _,r in leaves),rows[end][j],1e-7,'Mt')
    w.close()
    return dict(schemaVersion=1,scope=scope,sources=source_files,facts=facts,reconciliations=checks,excluded=excluded)
